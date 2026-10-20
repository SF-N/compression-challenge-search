"""Re-measure every submitted default-row configuration on the CURRENT official
challenge arrays with the official checker and compare with the sheet value.
Non-destructive: writes audit_submissions.jsonl (resumable) and a summary."""
import json, sys, time
from pathlib import Path
import numpy as np
from numcodecs.registry import get_codec
from openpyxl import load_workbook

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE))
from search import get_requirements
OUT = HERE / "audit_submissions.jsonl"
done = {}
if OUT.exists():
    for line in OUT.read_text().splitlines():
        c = json.loads(line); done[(c["lk"], c["var"])] = c

def log(m): print(f"[{time.strftime('%H:%M:%S')}] {m}", flush=True)

def main():
    from compression_requirement_checks import check_safety_requirements
    NS = {"nan": float("nan"), "inf": float("inf"), "__builtins__": {}}
    for lk, f in (("pressure", "ERA5-Pressure.xlsx"), ("single", "ERA5-Single.xlsx")):
        ws = load_workbook(HERE / "submissions" / f).active
        hdr = [c.value for c in ws[1]]
        iv, ic, idt, icfg = hdr.index("Variable"), hdr.index("Compression"), hdr.index("Default/all timestep(s)"), hdr.index("Configuration")
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[iv] is None or str(r[idt]) != "default": continue
            v = str(r[iv])
            if (lk, v) in done: continue
            t0 = time.time()
            cfg = eval(str(r[icfg]), NS)
            x = np.load(HERE / "data/era5" / f"{lk}__{v}.npz")["arr"]
            reqs = get_requirements(lk, v)
            try:
                codec = get_codec(cfg)
                e = codec.encode(x)
                d = np.asarray(codec.decode(np.frombuffer(bytes(e), np.uint8))).reshape(x.shape)
                ok = bool(check_safety_requirements(original=x, reconstructed=d, requirements=reqs))
                cr = x.nbytes / len(e)
                rec = dict(lk=lk, var=v, sheet_cr=float(r[ic]), measured_cr=cr, official_ok=ok,
                           rel_diff=cr / float(r[ic]) - 1, time_s=round(time.time() - t0, 1))
            except Exception as ex:
                rec = dict(lk=lk, var=v, sheet_cr=float(r[ic]), error=f"{type(ex).__name__}: {str(ex)[:150]}")
            with open(OUT, "a") as fh: fh.write(json.dumps(rec) + "\n")
            done[(lk, v)] = rec
            log(f"{lk} {v}: sheet {rec['sheet_cr']:.2f} measured {rec.get('measured_cr', float('nan')):.2f} ok={rec.get('official_ok')} {rec.get('error','')}")
    bad = [c for c in done.values() if c.get("error") or not c.get("official_ok") or abs(c.get("rel_diff", 0)) > 0.005]
    log(f"AUDIT DONE: {len(done)} rows, {len(bad)} discrepancies")
    for c in bad: log(f"   {c}")

if __name__ == "__main__":
    main()
